import io
import json
import os
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import requests

import main
from utils.parse_schedule import parse_schedule
from utils.sync_state import build_digest, stage_state
import webhook


class GeneratorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        cwd = Path.cwd()
        os.chdir(self.temp.name)
        self.addCleanup(os.chdir, cwd)
        env = patch.dict(os.environ, {}, clear=True)
        env.start()
        self.addCleanup(env.stop)
        output = patch("sys.stdout", new_callable=io.StringIO)
        output.start()
        self.addCleanup(output.stop)
        Path("index.html").write_text("<!-- CALENDARS -->\nOstatnia aktualizacja: <!-- UPDATED -->")
        self.combos = Path("combos.json").resolve()
        self.combos.write_text(json.dumps({
            "gomberman": ["i-rok-3-sem-5-gl1", "i-rok-3-sem-5-gl2"],
            "sztywne-gity": ["i-rok-3-sem-5-gl2"],
        }))
        for patcher in (patch("main.COMBOS_FILE", str(self.combos)),
                        patch("main.load_schedule"), patch("main.load_rooms", return_value={}),
                        patch("main.load_teachers", return_value={})):
            patcher.start()
            self.addCleanup(patcher.stop)
        self.write_excel()

    def write_excel(self, second_group=True, room="L1"):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet["A1"] = "I STOPIEŃ"
        sheet["C2"] = "ROK 3 sem 5"
        if second_group:
            sheet.merge_cells("C2:D2")
        sheet["A3"] = "sobota · Z1"
        sheet["C3"] = "GL1"
        sheet["A4"] = "03.10\n2026"
        sheet["C4"] = f"Sieci L\nJan Kowalski\n08:00–09:30 {room}"
        if second_group:
            sheet["D3"] = "GL2"
            sheet["D4"] = "Grafika L\nAnna Nowak\n10:00–11:30 L2"
        sheet["A6"] = "LEGENDA"
        sheet["A7"] = "Sieci"
        sheet["A8"] = "Grafika"
        book.save("plan.xlsx")
        book.close()

    def establish_checkpoint(self):
        main.main()
        stage_state("build", {}, "sync-state")
        os.environ["SYNC_STATE_DIR"] = "sync-state"
        os.environ["GITHUB_OUTPUT"] = str(Path("outputs.txt").resolve())

    def test_removed_column_removes_link_and_updates_combos_without_renaming_them(self):
        self.establish_checkpoint()
        old_uid = main.icalendar.Calendar.from_ical(Path("build/i-rok-3-sem-5-gl1.ics").read_bytes()).walk("VEVENT")[0]["UID"]
        self.write_excel(second_group=False, room="L4")
        main.main()
        files = {p.name for p in Path("build").glob("*.ics")}
        self.assertEqual({"i-rok-3-sem-5-gl1.ics", "gomberman.ics", "sztywne-gity.ics"}, files)
        page = Path("build/index.html").read_text()
        self.assertNotIn("i-rok-3-sem-5-gl2.ics", page)
        self.assertNotIn("gomberman.ics", page)
        current = main.icalendar.Calendar.from_ical(Path("build/i-rok-3-sem-5-gl1.ics").read_bytes()).walk("VEVENT")
        self.assertEqual(str(old_uid), str(current[0]["UID"]))
        self.assertEqual("L4", str(current[0]["LOCATION"]))
        combo = main.icalendar.Calendar.from_ical(Path("build/gomberman.ics").read_bytes()).walk("VEVENT")
        self.assertEqual(1, len(combo))
        self.assertIn("Grupa usunięta", webhook.format_message(webhook.dir_compare()))

    @patch("webhook.time.sleep")
    @patch("webhook.requests.post", side_effect=requests.Timeout)
    def test_failed_notification_preserves_baseline_and_is_retried_without_source_change(self, post, sleep):
        self.establish_checkpoint()
        completed = build_digest("sync-state/build")
        self.write_excel(second_group=False)
        main.main()
        os.environ["DISCORD_WEBHOOK_URL"] = "https://example.test/secret"
        with self.assertRaises(RuntimeError):
            webhook.main()
        self.assertEqual(completed, build_digest("sync-state/build"))
        # A new Actions runner starts with the same checkpoint, not next-state.
        shutil.rmtree("next-state")
        Path("outputs.txt").write_text("")
        main.main()
        self.assertIn("changed=true", Path("outputs.txt").read_text())
        self.assertIn("Grupa usunięta", webhook.format_message(webhook.dir_compare()))
        self.assertEqual(3, post.call_count)

    def test_unpublished_candidate_is_retried_after_deployment_failure(self):
        self.establish_checkpoint()
        self.write_excel(room="L4")
        main.main()
        completed = build_digest("sync-state/build")
        # Simulate termination before deploy confirmation/checkpoint upload.
        shutil.rmtree("next-state")
        Path("outputs.txt").write_text("")
        main.main()
        self.assertIn("changed=true", Path("outputs.txt").read_text())
        self.assertEqual(completed, build_digest("sync-state/build"))

    def test_successful_checkpoint_makes_next_unchanged_run_a_noop(self):
        self.establish_checkpoint()
        main.main()
        self.assertIn("changed=false", Path("outputs.txt").read_text())
        self.assertEqual([], webhook.dir_compare())

    def test_source_reverted_after_unfinished_deployment_must_still_be_published(self):
        self.establish_checkpoint()
        self.write_excel(room="L4")
        main.main()
        # The B build may be live, but notification/checkpoint failed. The next
        # restore detects its newer Pages artifact and supplies FORCE_DEPLOY.
        shutil.rmtree("next-state")
        self.write_excel(room="L1")
        os.environ["FORCE_DEPLOY"] = "true"
        Path("outputs.txt").write_text("")
        main.main()
        self.assertEqual(build_digest("sync-state/build"), build_digest("build"))
        self.assertIn("changed=true", Path("outputs.txt").read_text())

    def test_invalid_source_cannot_be_mistaken_for_removed_groups(self):
        self.establish_checkpoint()
        old = Path("build/i-rok-3-sem-5-gl1.ics").read_bytes()
        book = openpyxl.Workbook()
        book.save("plan.xlsx")
        book.close()
        with self.assertRaisesRegex(RuntimeError, "No day headers"):
            main.main()
        self.assertEqual(old, Path("build/i-rok-3-sem-5-gl1.ics").read_bytes())

    def test_event_in_unlabelled_column_is_rejected(self):
        book = openpyxl.load_workbook("plan.xlsx")
        book.active["D3"] = None
        book.save("plan.xlsx")
        book.close()
        with self.assertRaisesRegex(RuntimeError, "without a calendar"):
            parse_schedule("plan.xlsx")

    def test_invalid_time_is_rejected_before_timedelta_can_normalize_it(self):
        book = openpyxl.load_workbook("plan.xlsx")
        book.active["C4"] = "Sieci L\nJan Kowalski\n08:75–10:30 L1"
        book.save("plan.xlsx")
        book.close()
        with self.assertRaisesRegex(RuntimeError, "Invalid time"):
            parse_schedule("plan.xlsx")

    def test_later_day_cannot_silently_reassign_groups_by_column_position(self):
        book = openpyxl.load_workbook("plan.xlsx")
        sheet = book.active
        sheet.insert_rows(6, 3)
        sheet["A6"] = "niedziela · Z1"
        sheet["C6"] = "GL2"
        sheet["D6"] = "GL1"
        sheet["A7"] = "04.10\n2026"
        sheet["C7"] = "Sieci L\nJan Kowalski\n08:00–09:30 L1"
        sheet["D7"] = "Grafika L\nAnna Nowak\n10:00–11:30 L2"
        book.save("plan.xlsx")
        book.close()
        with self.assertRaisesRegex(RuntimeError, "calendar header"):
            parse_schedule("plan.xlsx")

    def test_new_language_column_without_allocated_events_can_be_published(self):
        book = openpyxl.load_workbook("plan.xlsx")
        book.active["D3"] = "JEZYK J1"
        book.active["D4"] = None
        book.save("plan.xlsx")
        book.close()
        main.main()
        self.assertTrue(Path("build/i-rok-3-sem-5-jezyk-j1.ics").exists())
        self.assertIn("J1", Path("build/index.html").read_text())

    def test_combo_deduplicates_a_shared_lecture_but_keeps_distinct_rooms(self):
        book = openpyxl.load_workbook("plan.xlsx")
        sheet = book.active
        sheet.merge_cells("C5:D5")
        sheet["C5"] = "Sieci W\nJan Kowalski\n12:00–13:30 L3"
        book.save("plan.xlsx")
        book.close()
        main.main()
        cal = main.icalendar.Calendar.from_ical(Path("build/gomberman.ics").read_bytes())
        self.assertEqual(3, len(cal.walk("VEVENT")))
        book = openpyxl.load_workbook("plan.xlsx")
        book.active.unmerge_cells("C5:D5")
        book.active["D5"] = "Sieci W\nJan Kowalski\n12:00–13:30 L4"
        book.save("plan.xlsx")
        book.close()
        main.main()
        cal = main.icalendar.Calendar.from_ical(Path("build/gomberman.ics").read_bytes())
        self.assertEqual(4, len(cal.walk("VEVENT")))
