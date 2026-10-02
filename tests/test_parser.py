import datetime
from pathlib import Path
import tempfile
import unittest

import openpyxl
from icalendar import Calendar

from main import calendar_events, new_calendar
from utils.parse_schedule import Event, Rubric, parse_schedule


class ParserTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'plan.xlsx'
        self.book = openpyxl.Workbook()
        self.addCleanup(self.book.close)
        self.sheet = self.book.active
        for cell, value in {'A1': 'I STOPIEŃ', 'C2': 'ROK 3 sem 5', 'A3': 'sobota · Z1',
                            'C3': 'GL1', 'D3': 'GL2', 'A4': '24.10\n2026',
                            'C4': 'Sieci L\nJan Kowalski\n08:00–09:30 L1',
                            'D4': 'Grafika L\nAnna Nowak\n10:00–11:30 L2',
                            'A9': 'LEGENDA', 'A10': 'Sieci', 'A11': 'Grafika'}.items():
            self.sheet[cell] = value
        self.sheet.merge_cells('C2:D2')

    def parse(self):
        self.book.save(self.path)
        return parse_schedule(self.path)

    def test_shared_merged_lecture_belongs_to_both_groups_once(self):
        self.sheet.merge_cells('C5:D6')
        self.sheet['C5'] = 'Sieci W\nJan Kowalski\n12:00–13:30 L3'
        rubrics = self.parse()
        self.assertEqual([2, 2], [len(r.events) for r in rubrics])
        self.assertIs(rubrics[0].events[-1], rubrics[1].events[-1])
        self.assertEqual('', rubrics[0].events[-1].groups)

    def test_language_merged_across_groups_is_split_once(self):
        self.sheet.merge_cells('C5:D5')
        self.sheet['C5'] = 'Język obcy angielski C · J1\nAnna Nowak\n12:00–13:30 ONLINE'
        rubrics = self.parse()
        languages = [r for r in rubrics if r.kind == 'language']
        self.assertEqual(1, len(languages))
        self.assertEqual('J1', languages[0].group)
        self.assertEqual(1, len(languages[0].events))
        self.assertTrue(all(len(r.events) == 1 for r in rubrics if r.kind == 'base'))

    def test_second_degree_has_its_own_date_and_group(self):
        for cell, value in {'F1': 'II STOPIEŃ', 'H2': 'ROK 1 sem 1', 'F3': 'sobota · Z1',
                            'H3': 'GL1', 'F4': '25.10\n2026',
                            'H4': 'Grafika P\nAnna Nowak\n12:00–13:30 L2'}.items():
            self.sheet[cell] = value
        rubrics = self.parse()
        second = [r for r in rubrics if r.degree == 'II'][0]
        self.assertEqual('ii-rok-1-sem-1-gl1', second.slug)
        self.assertEqual(datetime.date(2026, 10, 25), second.events[0].start.date())
        self.assertEqual('GL1', second.events[0].groups)

    def test_missing_header_in_second_section_is_controlled_failure(self):
        for cell, value in {'F1': 'II STOPIEŃ', 'H2': 'ROK 1 sem 1', 'F3': 'sobota · Z1',
                            'H3': 'GL1', 'F4': '24.10\n2026',
                            'H4': 'Grafika P\nAnna Nowak\n12:00–13:30 L2',
                            'A6': 'niedziela · Z1', 'C6': 'GL1', 'D6': 'GL2',
                            'A7': '25.10\n2026', 'H6': 'GL1', 'F7': '25.10\n2026'}.items():
            self.sheet[cell] = value
        with self.assertRaisesRegex(RuntimeError, 'Missing day header'):
            self.parse()

    def test_elective_name_with_polish_letters_keeps_existing_slug_rules(self):
        self.sheet['D3'] = 'Łączność i sieci\nK01'
        self.sheet['D4'] = 'Łączność i sieci L · K01\nAnna Nowak\n10:00–11:30 ONLINE'
        rubric = self.parse()[1]
        self.assertEqual('elective', rubric.kind)
        self.assertEqual('K01', rubric.group)
        self.assertEqual('i-rok-3-sem-5-lacznosc-i-sieci-k01', rubric.slug)
        self.assertEqual('Łączność i sieci', rubric.events[0].subject)

    def test_duplicate_normalized_calendar_names_are_rejected(self):
        self.sheet['D3'] = 'GL1'
        with self.assertRaisesRegex(RuntimeError, 'Duplicate calendar names'):
            self.parse()

    def test_missing_date_is_not_interpreted_as_removed_events(self):
        self.sheet['A4'] = ''
        with self.assertRaisesRegex(RuntimeError, 'Missing date'):
            self.parse()

    def test_timetable_survives_dst_change_with_same_local_hour(self):
        for cell, value in {'A6': 'niedziela · Z1', 'C6': 'GL1', 'D6': 'GL2',
                            'A7': '25.10\n2026',
                            'C7': 'Sieci L\nJan Kowalski\n08:00–09:30 L1'}.items():
            self.sheet[cell] = value
        rubric = self.parse()[0]
        cal = Calendar.from_ical(new_calendar('DST', calendar_events(rubric, {}, {})).to_ical())
        starts = [e['DTSTART'].dt for e in cal.walk('VEVENT')]
        self.assertEqual([8, 8], [v.hour for v in starts])
        self.assertEqual([7200, 3600], [v.utcoffset().total_seconds() for v in starts])
        self.assertEqual([6, 7], [v.astimezone(datetime.timezone.utc).hour for v in starts])

    def test_long_polish_subject_round_trips_and_online_location_is_preserved(self):
        subject = 'Zażółć gęślą jaźń; sieci, protokoły i bezpieczeństwo ' * 5
        event = Event(subject, 'L', 'GL1', 'Anna Nowak', 'ONLINE',
                      datetime.datetime(2026, 10, 3, 8), datetime.datetime(2026, 10, 3, 9, 30), 'Z1')
        rubric = Rubric('I', 'Rok 3 sem 5', 'GL1', 'i-rok-3-sem-5-gl1', events=[event])
        raw = new_calendar('Polski', calendar_events(rubric, {}, {})).to_ical()
        restored = Calendar.from_ical(raw).walk('VEVENT')[0]
        self.assertEqual('ONLINE – LAB ' + subject + ', grupa GL1', str(restored['SUMMARY']))
        self.assertEqual('Online', str(restored['LOCATION']))
        self.assertTrue(all(len(line) <= 75 for line in raw.split(b'\r\n')))
