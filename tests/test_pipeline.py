"""Exercise actual Python stages across fresh runners, with in-memory hosting/API.

GitHub's upload/deploy steps are simulated; no real service receives requests.
"""
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch
from urllib.parse import urlsplit
import zipfile

import openpyxl
from icalendar import Calendar

import main
from utils.sync_state import restore_state
import webhook


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        cwd = Path.cwd()
        self.addCleanup(os.chdir, cwd)
        self.artifacts, self.messages, self.site = [], [], {}
        self.runs, self.deployments = 0, 0
        self.source = self.make_source()
        self.env = patch.dict(os.environ, {}, clear=True)
        self.env.start()
        self.addCleanup(self.env.stop)
        output = patch('sys.stdout', new_callable=io.StringIO)
        self.output = output.start()
        self.addCleanup(output.stop)
        self.combos = {'gomberman': ['i-rok-3-sem-5-gl1', 'i-rok-3-sem-5-gl2'],
                       'sztywne-gity': ['i-rok-3-sem-5-gl2']}
        # Existing site A, as encountered during migration.
        runner = self.prepare_runner()
        with self.generator_dependencies(runner):
            main.main()
        self.site = self.read_directory(runner / 'build')

    def make_source(self, second=True, room='L1'):
        book = openpyxl.Workbook()
        sheet = book.active
        for cell, value in {'A1': 'I STOPIEŃ', 'C2': 'ROK 3 sem 5', 'A3': 'sobota · Z1',
                            'C3': 'GL1', 'A4': '03.10\n2026',
                            'C4': f'Sieci L\nJan Kowalski\n08:00–09:30 {room}',
                            'A6': 'LEGENDA', 'A7': 'Sieci', 'A8': 'Grafika'}.items():
            sheet[cell] = value
        if second:
            sheet.merge_cells('C2:D2')
            sheet['D3'] = 'GL2'
            sheet['D4'] = 'Grafika L\nAnna Nowak\n10:00–11:30 L2'
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        return stream.getvalue()

    def prepare_runner(self):
        self.runs += 1
        runner = self.root / str(self.runs)
        (runner / 'data').mkdir(parents=True)
        (runner / 'data/combos.json').write_text(json.dumps(self.combos))
        (runner / 'index.html').write_text('<!-- CALENDARS -->\nOstatnia aktualizacja: <!-- UPDATED -->')
        os.chdir(runner)
        return runner

    def generator_dependencies(self, runner):
        from contextlib import ExitStack
        stack = ExitStack()
        stack.enter_context(patch('main.COMBOS_FILE', str(runner / 'data/combos.json')))
        stack.enter_context(patch('main.load_schedule', side_effect=lambda: Path('plan.xlsx').write_bytes(self.source)))
        stack.enter_context(patch('main.load_rooms', return_value={}))
        stack.enter_context(patch('main.load_teachers', return_value={}))
        return stack

    @staticmethod
    def read_directory(directory):
        return {str(path.relative_to(directory)): path.read_bytes()
                for path in directory.rglob('*') if path.is_file()}

    def upload(self, name, directory=None):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w') as archive:
            for name_in_archive, data in self.read_directory(directory).items() if directory else []:
                archive.writestr(name_in_archive, data)
        identity = max((a['id'] for a, _ in self.artifacts), default=0) + 1
        self.artifacts.append(({
            'name': name, 'id': identity, 'expired': False,
            'created_at': f'2026-10-01T00:00:{identity:02d}Z',
            'archive_download_url': f'https://api.github.com/archive/{identity}',
            'workflow_run': {'head_branch': 'main', 'repository_id': 1, 'head_repository_id': 1},
        }, stream.getvalue()))

    def get(self, session, url):
        path = urlsplit(url).path
        if path.endswith('/actions/artifacts'):
            return Mock(json=Mock(return_value={'artifacts': [a for a, _ in reversed(self.artifacts)]}))
        if path.startswith('/archive/'):
            identity = int(path.rsplit('/', 1)[1])
            return Mock(content=next(data for metadata, data in self.artifacts if metadata['id'] == identity))
        if urlsplit(url).netloc == 'planpk.linguin.dev':
            return Mock(content=self.site[path.lstrip('/')])
        raise AssertionError(f'Unexpected network request: {path}')

    def sync(self, failure=None):
        runner = self.prepare_runner()
        outputs = runner / 'outputs.txt'
        with patch.dict(os.environ, {
            'SYNC_STATE_DIR': 'sync-state', 'GITHUB_TOKEN': 'test', 'GITHUB_REPOSITORY': 'owner/repo',
            'GITHUB_OUTPUT': str(outputs), 'DISCORD_WEBHOOK_URL': 'https://example.test/not-real',
        }), patch('utils.sync_state._get', side_effect=self.get), self.generator_dependencies(runner):
            restore_state()
            values = dict(line.split('=', 1) for line in outputs.read_text().splitlines())
            if values['bootstrap'] == 'true':
                self.upload(f'sync-state-initial-{self.runs}-1', runner / 'sync-state')
            os.environ['FORCE_DEPLOY'] = values['force_deploy']
            main.main()
            values = dict(line.split('=', 1) for line in outputs.read_text().splitlines())
            if values['changed'] == 'true':
                self.upload(f'github-pages-{self.runs}-1')
                if failure == 'deploy':
                    raise RuntimeError('simulated deployment failure')
                self.site = self.read_directory(runner / 'build')
                self.deployments += 1

            def post(url, **kwargs):
                if failure == 'discord':
                    return Mock(status_code=503, headers={})
                self.messages.append(kwargs['json']['content'])
                return Mock(status_code=200, json=Mock(return_value={'id': str(len(self.messages))}))

            with patch('webhook.requests.post', side_effect=post), patch('webhook.time.sleep'):
                webhook.main()
            if failure == 'checkpoint':
                raise RuntimeError('simulated checkpoint upload failure')
            self.upload(f'sync-state-{self.runs}-1', runner / 'next-state')
        return values['changed'] == 'true'

    def test_bootstrap_then_unchanged_runs_refresh_state_without_republishing(self):
        self.assertTrue(self.sync())
        self.assertFalse(self.sync())
        self.assertFalse(self.sync())
        self.assertEqual(1, self.deployments)
        self.assertEqual([], self.messages)
        self.assertEqual(4, sum(a['name'].startswith('sync-state-') for a, _ in self.artifacts))

    def test_removal_survives_failed_deploy_then_stops_notifying_after_success(self):
        self.sync()
        original_site = dict(self.site)
        self.source = self.make_source(second=False)
        with self.assertRaisesRegex(RuntimeError, 'deployment failure'):
            self.sync(failure='deploy')
        self.assertEqual(original_site, self.site)
        self.assertEqual([], self.messages)
        self.assertTrue(self.sync())
        self.assertNotIn('i-rok-3-sem-5-gl2.ics', self.site)
        self.assertIn('gomberman.ics', self.site)
        self.assertIn('sztywne-gity.ics', self.site)
        self.assertIn('Grupa usunięta', self.messages[0])
        self.assertFalse(self.sync())
        self.assertEqual(1, len(self.messages))

    def test_expired_checkpoint_recovers_then_retries_failed_notification(self):
        self.sync()
        self.artifacts[-1][0]['expired'] = True
        self.source = self.make_source(room='L4')
        with self.assertRaisesRegex(RuntimeError, 'HTTP 503'):
            self.sync(failure='discord')
        self.assertIn('::warning::Ostatni zapis synchronizacji wygasł.', self.output.getvalue())
        self.assertEqual([], self.messages)
        self.assertTrue(self.sync())
        self.assertIn('Zmiana sali: L1 -> L4', self.messages[0])
        self.assertFalse(self.sync())
        self.assertEqual(1, len(self.messages))

    def test_lost_checkpoint_after_deploy_recovers_and_notifies_future_changes(self):
        self.sync()
        self.source = self.make_source(room='L4')
        with self.assertRaisesRegex(RuntimeError, 'HTTP 503'):
            self.sync(failure='discord')
        self.artifacts = [(metadata, data) for metadata, data in self.artifacts
                          if not metadata['name'].startswith('sync-state-')]
        self.assertTrue(self.sync())
        self.assertIn('::warning::Brak zapisu synchronizacji.', self.output.getvalue())
        # The already-published change cannot be reconstructed after losing state.
        self.assertEqual([], self.messages)
        self.source = self.make_source(room='L5')
        self.assertTrue(self.sync())
        self.assertIn('Zmiana sali: L4 -> L5', self.messages[0])
        self.assertFalse(self.sync())
        self.assertEqual(1, len(self.messages))

    def test_recovery_keeps_published_cohort_semester(self):
        self.sync()
        self.update_source_cells({'C2': 'ROK 4 sem 7', 'A4': '02.10\n2027'})
        self.sync()
        sent = len(self.messages)
        self.artifacts.clear()
        self.assertTrue(self.sync())
        self.assertEqual(sent, len(self.messages))
        self.update_source_cells({'C4': 'Sieci L\nJan Kowalski\n08:00–09:30 L4'})
        self.assertTrue(self.sync())
        self.assertIn('Rok 4 sem 7', self.messages[-1])
        self.assertIn('Zmiana sali: L1 -> L4', self.messages[-1])
        self.assertFalse(self.sync())

    def test_removal_notification_survives_successful_deploy_and_failed_discord(self):
        self.sync()
        self.source = self.make_source(second=False)
        with self.assertRaisesRegex(RuntimeError, 'HTTP 503'):
            self.sync(failure='discord')
        self.assertNotIn('i-rok-3-sem-5-gl2.ics', self.site)
        self.assertTrue(self.sync())
        self.assertIn('Grupa usunięta', self.messages[0])
        self.assertFalse(self.sync())
        self.assertEqual(1, len(self.messages))

    def test_source_revert_after_discord_failure_restores_actual_site(self):
        self.sync()
        self.source = self.make_source(room='L4')
        with self.assertRaisesRegex(RuntimeError, 'HTTP 503'):
            self.sync(failure='discord')
        self.assertIn(b'LOCATION:L4', self.site['i-rok-3-sem-5-gl1.ics'])
        self.source = self.make_source(room='L1')
        self.assertTrue(self.sync())
        self.assertIn(b'LOCATION:L1', self.site['i-rok-3-sem-5-gl1.ics'])
        self.assertFalse(self.sync())
        self.assertEqual([], self.messages)

    def test_failed_checkpoint_can_duplicate_a_message_but_cannot_lose_it(self):
        self.sync()
        self.source = self.make_source(room='L4')
        with self.assertRaisesRegex(RuntimeError, 'checkpoint upload'):
            self.sync(failure='checkpoint')
        self.assertEqual(1, len(self.messages))
        self.sync()
        self.assertEqual([self.messages[0]] * 2, self.messages)
        self.assertFalse(self.sync())
        self.assertEqual(2, len(self.messages))

    def test_invalid_excel_never_deploys_or_announces_mass_deletion(self):
        self.sync()
        previous = dict(self.site)
        previous_artifacts = len(self.artifacts)
        self.source = b'PKnot-a-workbook'
        with self.assertRaises(zipfile.BadZipFile):
            self.sync()
        self.assertEqual(previous, self.site)
        self.assertEqual(previous_artifacts, len(self.artifacts))
        self.assertEqual([], self.messages)

    def update_source_cells(self, values):
        book = openpyxl.load_workbook(io.BytesIO(self.source))
        for cell, value in values.items():
            book.active[cell] = value
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        self.source = stream.getvalue()

    def site_events(self, name):
        return Calendar.from_ical(self.site[name]).walk('VEVENT')

    def test_full_feed_edit_cancel_last_class_and_add_new_class(self):
        self.sync()
        group = 'i-rok-3-sem-5-gl1.ics'
        unchanged_group = 'i-rok-3-sem-5-gl2.ics'
        other_uid = str(self.site_events(unchanged_group)[0]['UID'])
        self.update_source_cells({'C4': 'Sieci komputerowe L\nAnna Nowak\n10:00–12:00 L4'})
        self.assertTrue(self.sync())
        edited = self.site_events(group)
        self.assertEqual(1, len(edited))
        self.assertEqual('LAB Sieci komputerowe, grupa GL1', str(edited[0]['SUMMARY']))
        self.assertEqual((10, 12), (edited[0]['DTSTART'].dt.hour, edited[0]['DTEND'].dt.hour))
        self.assertEqual('L4', str(edited[0]['LOCATION']))
        self.assertIn('Anna Nowak', str(edited[0]['DESCRIPTION']))
        self.assertFalse(any(str(e['X-PK-KEY']) == 'LAB Sieci' for e in self.site_events('gomberman.ics')))

        # The column still exists, but all its classes have been cancelled.
        self.update_source_cells({'C4': None})
        self.assertTrue(self.sync())
        self.assertEqual([], self.site_events(group))
        self.assertIn(group, {entry['file'] for entry in json.loads(self.site['calendars.json'])})
        self.assertEqual(1, len(self.site_events('gomberman.ics')))
        self.assertIn('Wydarzenie usunięte', self.messages[-1])
        self.assertNotIn('Grupa usunięta', self.messages[-1])

        self.update_source_cells({'C4': 'Nowy przedmiot C\nJan Kowalski\n14:00–15:30 ONLINE'})
        self.assertTrue(self.sync())
        added = self.site_events(group)
        self.assertEqual(1, len(added))
        self.assertEqual('‼️ONLINE‼️ – ĆWICZENIA Nowy przedmiot, grupa GL1', str(added[0]['SUMMARY']))
        self.assertEqual('Online', str(added[0]['LOCATION']))
        self.assertEqual((14, 15, 30), (added[0]['DTSTART'].dt.hour, added[0]['DTEND'].dt.hour, added[0]['DTEND'].dt.minute))
        self.assertEqual(2, len(self.site_events('gomberman.ics')))
        self.assertEqual(other_uid, str(self.site_events(unchanged_group)[0]['UID']))
        self.assertIn('Nowe wydarzenie', self.messages[-1])
        self.assertFalse(self.sync())

    def test_simultaneous_same_subject_entries_are_published_in_group_and_combo(self):
        self.sync()
        self.update_source_cells({'C5': 'Sieci L · GL3\nAnna Nowak\n08:00–09:30 L4'})
        self.assertTrue(self.sync())
        group = self.site_events('i-rok-3-sem-5-gl1.ics')
        self.assertEqual(2, len(group))
        self.assertEqual({'L1', 'L4'}, {str(e['LOCATION']) for e in group})
        self.assertEqual(2, len({str(e['UID']) for e in group}))
        self.assertEqual(3, len(self.site_events('gomberman.ics')))
        self.assertFalse(self.sync())

    def test_cohort_progresses_to_eight_and_only_commits_after_notification(self):
        def completed_semester():
            content = next(data for metadata, data in reversed(self.artifacts)
                           if metadata['name'].startswith('sync-state-'))
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                return json.loads(archive.read('state.json'))['notification_semester']

        self.sync()
        self.assertEqual(5, completed_semester())
        for semester, day in ((6, '27.02\n2027'), (7, '02.10\n2027'), (8, '26.02\n2028')):
            self.update_source_cells({'C2': f'ROK {(semester + 1) // 2} sem {semester}', 'A4': day})
            if semester == 6:
                with self.assertRaisesRegex(RuntimeError, 'HTTP 503'):
                    self.sync(failure='discord')
                self.assertEqual(5, completed_semester())
            self.assertTrue(self.sync())
            self.assertEqual(semester, completed_semester())
            self.assertIn(f'Rok {(semester + 1) // 2} sem {semester}', self.messages[-1])
            self.assertNotIn('Grupa usunięta', self.messages[-1])
            self.assertTrue(self.messages[-1].startswith('<@&1286988227617488896>'))
            self.assertFalse(self.sync())
        self.update_source_cells({'C2': 'ROK 3 sem 5', 'A4': '07.10\n2028'})
        self.sync()
        self.assertEqual(8, completed_semester())
        sent = len(self.messages)
        self.update_source_cells({'C4': 'Sieci L\nJan Kowalski\n08:00–09:30 L4'})
        self.sync()
        self.assertEqual(sent, len(self.messages))
        self.assertEqual(8, completed_semester())
