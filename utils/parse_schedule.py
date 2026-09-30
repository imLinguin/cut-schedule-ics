import datetime
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field

import openpyxl

"""
Layout of the Excel exported by FK Planer (plan-ns/download.php):

- The sheet is split into sections side by side (I stopień on the left, II stopień on the right).
  Each section starts with a date column followed by a time column.
- Row above the first day header holds year headers ("ROK 3 sem 5"), merged across their columns.
- Every day starts with a header row ("sobota · Z1") that repeats the column labels
  ("GL1", "Programowanie na platformie .NET\nK01", ...). Each column is one "rubric".
- The row below a day header holds the date ("03.10\n2026").
- Events are (merged) cells in the form "Subject ACT[ · groups]\nTeacher\nHH:MM–HH:MM ROOM".
  A cell merged across columns belongs to every rubric it covers.
- Below the table there is a legend listing every subject name.

Rubric kinds:
- base: groups of a year (GL1, "DS GL1")
- elective: an elective subject group ("Programowanie na platformie .NET K01")
- language: a language group ("JEZYK J1" columns, and "Język obcy" groups which the
  Excel puts in GL columns but which students pick on their own)
"""

# Excel activity letter -> (event title prefix, category)
ACTIVITIES = {
    "W": ("WYKŁAD", "wykład"),
    "C": ("ĆW", "ćwiczenia"),
    "L": ("LAB", "laboratorium"),
    "P": ("PROJEKT", "projekt"),
    "S": ("SEMINARIUM", "seminarium"),
}

DAY_HEADER = re.compile(r"^\S+ · (Z\d+)$")
DATE = re.compile(r"^(\d\d)\.(\d\d)\s+(\d{4})$")
TIME = re.compile(r"^(\d\d):(\d\d)[–-](\d\d):(\d\d)(?: (.*))?$")
TITLE = re.compile(r"^(?P<subject>.+?) (?P<activity>[WCLPS])(?:(?: ·)? (?P<groups>.+))?$")
ELECTIVE = re.compile(r"^(?P<subject>.+) (?P<group>[KPĆ]\d+)$")
LANGUAGE_COLUMN = re.compile(r"\b(?P<group>J\d+)$")
BASE_GROUP = re.compile(r"\bGL\d+$")
LANGUAGE_SUBJECT = re.compile(r"^Język obcy\b")


@dataclass
class Event:
    subject: str
    activity: str
    groups: str
    teacher: str
    room: str
    start: datetime.datetime
    end: datetime.datetime
    weekend: str


@dataclass
class Rubric:
    degree: str
    year: str
    label: str
    slug: str
    kind: str = "base"
    subject: str = None
    group: str = None
    events: list = field(default_factory=list)


def slugify(value: str) -> str:
    value = value.replace("ł", "l").replace("Ł", "L")
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def _text(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _merged_lookup(ws):
    # (row, col) of the top-left cell -> (last row, first col, last col)
    lookup = {}
    covered = set()
    for rg in ws.merged_cells.ranges:
        lookup[(rg.min_row, rg.min_col)] = (rg.max_row, rg.min_col, rg.max_col)
        for row in range(rg.min_row, rg.max_row + 1):
            for col in range(rg.min_col, rg.max_col + 1):
                if (row, col) != (rg.min_row, rg.min_col):
                    covered.add((row, col))
    return lookup, covered


def _merged_value(ws, merged, row, col):
    for (top, left), (bottom, _, right) in merged.items():
        if top <= row <= bottom and left <= col <= right:
            return ws.cell(top, left).value
    return ws.cell(row, col).value


def _find_sections(ws, header_row):
    # A section starts wherever the day header text ("sobota · Z1") is found
    starts = [
        col
        for col in range(1, ws.max_column + 1)
        if DAY_HEADER.match(_text(ws.cell(header_row, col).value))
    ]
    sections = []
    for i, start in enumerate(starts):
        end = starts[i + 1] - 1 if i + 1 < len(starts) else ws.max_column
        degree = None
        for row in range(1, header_row):
            match = re.search(r"\b(I+) STOPIEŃ", _text(ws.cell(row, start).value).upper())
            if match:
                degree = match.group(1)
                break
        if not degree:
            raise RuntimeError(f"No degree title above column {start}")
        sections.append((degree, start, end))
    return sections


def _legend_subjects(ws, legend_row):
    subjects = []
    for row in range(legend_row + 1, ws.max_row + 1):
        value = _text(ws.cell(row, 1).value)
        if value:
            subjects.append(value)
    # Longest first, so "Metody Obliczeniowe w Nauce i Technice" wins over shorter prefixes
    return sorted(subjects, key=len, reverse=True)


def _parse_title(title: str, subjects: list):
    for subject in subjects:
        if title.startswith(subject + " "):
            rest = title[len(subject) + 1 :]
            activity, _, groups = rest.partition(" ")
            if activity in ACTIVITIES:
                return subject, activity, groups.lstrip("· ").strip()
    match = TITLE.match(title)
    if match:
        return match["subject"], match["activity"], (match["groups"] or "").strip()
    return title, "", ""


def _split_languages(rubrics: list) -> list:
    """Moves "Język obcy" groups out of GL calendars into their own calendars."""
    languages = {}
    for rubric in rubrics:
        if rubric.kind != "base":
            continue
        kept = []
        for event in rubric.events:
            if not LANGUAGE_SUBJECT.match(event.subject):
                kept.append(event)
                continue
            key = (rubric.degree, rubric.year, event.subject, event.groups)
            if key not in languages:
                label = f"{event.subject} {event.groups}".strip()
                languages[key] = Rubric(
                    rubric.degree,
                    rubric.year,
                    label,
                    slugify(f"{rubric.degree} {rubric.year} {label}"),
                    kind="language",
                    subject=event.subject,
                    group=event.groups,
                )
            # The same cell can span several GL columns
            if not any(event is known for known in languages[key].events):
                languages[key].events.append(event)
        rubric.events = kept
    return list(languages.values())


def parse_schedule(path: str) -> list:
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    merged, covered = _merged_lookup(ws)

    day_rows = [
        row
        for row in range(1, ws.max_row + 1)
        if DAY_HEADER.match(_text(ws.cell(row, 1).value))
    ]
    if not day_rows:
        raise RuntimeError("No day headers found, the Excel layout has changed")
    legend_row = next(
        (
            row
            for row in range(day_rows[-1], ws.max_row + 1)
            if "LEGENDA" in _text(ws.cell(row, 1).value).upper()
        ),
        ws.max_row + 1,
    )
    subjects = _legend_subjects(ws, legend_row)
    header_row = day_rows[0]
    year_row = header_row - 1

    rubrics = []
    column_rubrics = {}
    for degree, start, end in _find_sections(ws, header_row):
        for col in range(start + 2, end + 1):
            year = _text(_merged_value(ws, merged, year_row, col))
            label = _text(ws.cell(header_row, col).value)
            if not label or not year:
                continue
            year = re.sub(r"^ROK\b", "Rok", year)
            # "Rok 1 sem 1 DS" + "DS GL1" -> "...-ds-gl1" instead of "...-ds-ds-gl1"
            slug_label = label
            if label.split(" ")[0].lower() == year.split(" ")[-1].lower():
                slug_label = label.split(" ", 1)[1]
            slug = slugify(f"{degree} {year} {slug_label}")
            rubric = Rubric(degree, year, label, slug)
            elective = ELECTIVE.match(label)
            language = LANGUAGE_COLUMN.search(label)
            if elective:
                rubric.kind, rubric.subject, rubric.group = "elective", elective["subject"], elective["group"]
            elif language:
                rubric.kind, rubric.group = "language", language["group"]
            rubrics.append(rubric)
            column_rubrics[col] = rubric

        # Walk the days of this section
        for i, day_row in enumerate(day_rows):
            next_day = day_rows[i + 1] if i + 1 < len(day_rows) else legend_row
            weekend = DAY_HEADER.match(_text(ws.cell(day_row, start).value)).group(1)
            date_match = DATE.match(_text(ws.cell(day_row + 1, start).value))
            if not date_match:
                raise RuntimeError(f"Missing date below row {day_row}")
            day, month, year_no = map(int, date_match.groups())
            date = datetime.datetime(year_no, month, day)

            for row in range(day_row + 1, next_day):
                for col in range(start + 2, end + 1):
                    if (row, col) in covered:
                        continue
                    value = ws.cell(row, col).value
                    if not isinstance(value, str) or not value.strip():
                        continue
                    lines = [line.strip() for line in value.strip().split("\n")]
                    time_match = TIME.match(lines[-1])
                    if not time_match:
                        raise RuntimeError(
                            f"Unexpected cell {ws.cell(row, col).coordinate}: {value!r}"
                        )
                    h1, m1, h2, m2 = map(int, time_match.groups()[:4])
                    subject, activity, groups = _parse_title(lines[0], subjects)
                    _, first_col, last_col = merged.get((row, col), (row, col, col))
                    event_rubrics = [
                        column_rubrics[event_col]
                        for event_col in range(first_col, last_col + 1)
                        if event_col in column_rubrics
                    ]
                    # II stopień cells have no group for labs and projects, the column says it
                    column_groups = [BASE_GROUP.search(r.label) for r in event_rubrics]
                    if not groups and activity != "W" and column_groups and all(column_groups):
                        groups = "+".join(match[0] for match in column_groups)
                    event = Event(
                        subject=subject,
                        activity=activity,
                        groups=groups,
                        teacher=" ".join(lines[1:-1]),
                        room=(time_match.group(5) or "").strip(),
                        start=date + datetime.timedelta(hours=h1, minutes=m1),
                        end=date + datetime.timedelta(hours=h2, minutes=m2),
                        weekend=weekend,
                    )
                    for rubric in event_rubrics:
                        rubric.events.append(event)

    rubrics += _split_languages(rubrics)
    for rubric in rubrics:
        if rubric.kind == "language" and not rubric.subject and rubric.events:
            subjects_used = Counter(event.subject for event in rubric.events)
            rubric.subject = subjects_used.most_common(1)[0][0]

    slugs = [rubric.slug for rubric in rubrics]
    duplicates = {slug for slug in slugs if slugs.count(slug) > 1}
    if duplicates:
        raise RuntimeError(f"Duplicate calendar names: {duplicates}")
    return rubrics
